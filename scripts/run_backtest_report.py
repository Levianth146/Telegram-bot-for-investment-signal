"""Báo cáo backtest thuyết trình — bảng terminal + Excel/CSV.

Cột: B0 buy&hold | B1 TA | B2 CANSLIM | Framework (WF OOS) | VN-Index | VN30.

Final VN100 (6m OOS, daily, fundamentals, warmup 3y):
  python scripts/run_backtest_report.py --universe vn100 --with-fundamentals \\
      --signal-every 1 --oos-start 2025-03-22 --oos-end 2025-09-22 --warmup-years 3 \\
      --no-walk-forward --out-json store/backtest_final_vn100.json

  # Lần 2 (cache hit OHLCV + scoring_schedule theo năm):
  # cùng lệnh — year_*.pkl dưới data/cache/scoring_schedule/

  # Force refresh BCTC / giá:
  #   ... --refresh-fundamentals
  #   ... --refresh-data

Smoke nhanh (ít mã, cửa sổ ngắn — chứng minh bảng in được):
  python scripts/run_backtest_report.py --tickers FPT,VNM,GAS --start-date 2022-01-01 \\
      --oos-start 2022-01-01 --signal-every 42 --no-fundamentals \\
      --out-xlsx store/backtest_report_smoke.xlsx

Profile smoke:
  python scripts/profile_backtest_smoke.py --out outputs/performance/profile_after.txt
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from datetime import date
from pathlib import Path
from typing import Any, Mapping


def _safe_print(text: str) -> None:
    """In Unicode an toàn trên Windows cp1252."""
    try:
        print(text, flush=True)
    except UnicodeEncodeError:
        enc = getattr(sys.stdout, "encoding", None) or "utf-8"
        print(text.encode(enc, errors="replace").decode(enc, errors="replace"), flush=True)


def _fmt(val: Any, pct: bool = False, digits: int = 2) -> str:
    if val is None:
        return "—"
    try:
        x = float(val)
    except (TypeError, ValueError):
        return "—"
    if x != x:  # NaN
        return "—"
    if abs(x) == float("inf"):
        return "inf" if x > 0 else "-inf"
    if pct:
        return f"{x * 100:.{digits}f}"
    return f"{x:.{digits}f}"


def _shift_years(iso_date: str, years: int) -> str:
    """Lùi/tiến năm lịch (xấp xỉ warmup) giữ tháng-ngày."""
    y, m, d = int(iso_date[:4]), int(iso_date[5:7]), int(iso_date[8:10])
    try:
        return date(y + years, m, d).isoformat()
    except ValueError:
        return date(y + years, m, min(d, 28)).isoformat()


def _avg_return_per_trade(trades: list[dict]) -> float | None:
    closed = [t for t in trades if t.get("pnl") is not None]
    if not closed:
        return None
    rets: list[float] = []
    for t in closed:
        gr = t.get("gross_return")
        if gr is not None:
            try:
                rets.append(float(gr))
                continue
            except (TypeError, ValueError):
                pass
        entry = t.get("entry_price")
        exit_px = t.get("exit_price")
        try:
            if entry and exit_px and float(entry) > 0:
                rets.append(float(exit_px) / float(entry) - 1.0)
        except (TypeError, ValueError):
            continue
    if not rets:
        return None
    return float(sum(rets) / len(rets))


def _total_return(equity_curve: list[dict]) -> float | None:
    if not equity_curve:
        return None
    try:
        first = float(equity_curve[0]["equity"])
        last = float(equity_curve[-1]["equity"])
    except (TypeError, ValueError, KeyError, IndexError):
        return None
    if first <= 0:
        return None
    return last / first - 1.0


def _slice_equity_to_oos(
    result: Mapping[str, Any], oos_start: str, oos_end: str
) -> dict[str, Any]:
    """Cắt equity/trades về cửa sổ OOS (model đã warm-up trên full history)."""
    curve = [
        p
        for p in (result.get("equity_curve") or [])
        if oos_start <= str(p.get("date") or "") <= oos_end
    ]
    if curve:
        base = float(curve[0]["equity"]) or 1.0
        curve = [{**p, "equity": float(p["equity"]) / base} for p in curve]
    trades = [
        t
        for t in (result.get("trades") or [])
        if oos_start
        <= str(t.get("exit_date") or t.get("entry_date") or "")
        <= oos_end
    ]
    return {
        **dict(result),
        "equity_curve": curve,
        "trades": trades,
    }


def _column_from_result(
    label: str,
    result: Mapping[str, Any] | None,
    *,
    is_benchmark: bool = False,
    fold_meta: Mapping[str, Any] | None = None,
) -> dict[str, Any]:
    """Chuẩn hoá một cột báo cáo từ kết quả engine/baseline/WF."""
    empty = {
        "label": label,
        "n_trades": None,
        "win_rate": None,
        "avg_return_per_trade": None,
        "total_return": None,
        "max_drawdown": None,
        "profit_factor": None,
        "sharpe": None,
        "cagr": None,
        "sortino": None,
        "calmar": None,
        "n_folds": None,
        "fold_sharpe_std": None,
        "fold_sharpe_mean": None,
        "fold_sharpes": [],
        "gross_total_return": None,
        "cost_drag": None,
        "sharpe_bull_display": None,
        "sharpe_bear_display": None,
        "avg_exposure": None,
    }
    if not result:
        return empty
    m = dict(result.get("metrics") or {})
    trades = list(result.get("trades") or [])
    curve = list(result.get("equity_curve") or [])
    n_trades = m.get("n_trades")
    if is_benchmark:
        n_trades = None
        win = None
        avg_rt = None
        pf = None
    else:
        win = m.get("win_rate")
        avg_rt = _avg_return_per_trade(trades)
        pf = m.get("profit_factor")
        if not trades and n_trades is not None and int(n_trades or 0) > 0:
            win = None
            avg_rt = None
            pf = None
    out = {
        "label": label,
        "n_trades": n_trades,
        "win_rate": win,
        "avg_return_per_trade": avg_rt,
        "total_return": (
            m.get("total_return")
            if m.get("total_return") is not None
            else _total_return(curve)
        ),
        "max_drawdown": m.get("max_drawdown"),
        "profit_factor": pf,
        "sharpe": m.get("sharpe"),
        "cagr": m.get("cagr"),
        "sortino": m.get("sortino"),
        "calmar": m.get("calmar"),
        "n_folds": None,
        "fold_sharpe_std": None,
        "fold_sharpe_mean": None,
        "fold_sharpes": [],
        "gross_total_return": m.get("gross_total_return"),
        "cost_drag": m.get("cost_drag"),
        "sharpe_bull_display": m.get("sharpe_bull_display") or "N/A",
        "sharpe_bear_display": m.get("sharpe_bear_display") or "N/A",
        "avg_exposure": m.get("avg_exposure"),
    }
    if fold_meta:
        out["n_folds"] = fold_meta.get("n_folds")
        out["fold_sharpe_std"] = fold_meta.get("fold_sharpe_std")
        out["fold_sharpe_mean"] = fold_meta.get("fold_sharpe_mean")
        out["fold_sharpes"] = list(fold_meta.get("fold_sharpes") or [])
    return out


def _load_universe_csv(path: Path) -> list[str]:
    tickers: list[str] = []
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            raw = line.strip()
            if not raw or raw.startswith("#") or raw.lower() == "ticker":
                continue
            tickers.append(raw.split(",")[0].strip().upper())
    return tickers


def _load_smoke_tickers(config: Mapping[str, Any], root: Path) -> list[str]:
    uni = dict(config.get("universe") or {})
    rel = uni.get("smoke_file") or "data/universe/hose_liquid_35.csv"
    path = root / str(rel)
    if not path.is_file():
        raise SystemExit(f"smoke_file not found: {path}")
    return _load_universe_csv(path)


def _load_vn100_tickers(config: Mapping[str, Any], root: Path) -> list[str]:
    uni = dict(config.get("universe") or {})
    rel = uni.get("vn100_file") or "data/universe/vn100.csv"
    path = root / str(rel)
    if not path.is_file():
        raise SystemExit(f"vn100 file not found: {path}")
    return _load_universe_csv(path)


def _print_table(columns: list[dict[str, Any]]) -> None:
    headers = [c["label"] for c in columns]
    rows_spec = [
        ("So lenh", "n_trades", False, False),
        ("Win rate (%)", "win_rate", True, False),
        ("Return TB/lenh (%)", "avg_return_per_trade", True, False),
        ("Total Return (%)", "total_return", True, False),
        ("Max Drawdown (%)", "max_drawdown", True, False),
        ("Profit Factor", "profit_factor", False, False),
        ("Sharpe", "sharpe", False, False),
        ("CAGR (%)", "cagr", True, False),
        ("Sortino", "sortino", False, False),
        ("Calmar", "calmar", False, False),
        ("Avg exposure", "avg_exposure", True, False),
        ("Cost drag (%)", "cost_drag", True, False),
        ("n_folds", "n_folds", False, False),
        ("fold Sharpe mean", "fold_sharpe_mean", False, False),
        ("fold Sharpe std", "fold_sharpe_std", False, False),
    ]
    col0_w = 22
    col_w = 12
    line = "Chi tieu".ljust(col0_w) + "".join(h[:col_w].center(col_w) for h in headers)
    _safe_print(line)
    _safe_print("-" * len(line))
    for label, key, as_pct, _ in rows_spec:
        cells = []
        for col in columns:
            val = col.get(key)
            if key == "n_trades" or key == "n_folds":
                cells.append("—" if val is None else str(int(val)))
            else:
                cells.append(_fmt(val, pct=as_pct))
        _safe_print(label.ljust(col0_w) + "".join(c.center(col_w) for c in cells))


def _export_xlsx_or_csv(
    columns: list[dict[str, Any]],
    out_path: Path,
    meta: Mapping[str, Any],
) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    metric_rows = [
        ("Số lệnh", "n_trades", False),
        ("Win rate (%)", "win_rate", True),
        ("Return TB/lệnh (%)", "avg_return_per_trade", True),
        ("Total Return (%)", "total_return", True),
        ("Max Drawdown (%)", "max_drawdown", True),
        ("Profit Factor", "profit_factor", False),
        ("Sharpe", "sharpe", False),
        ("CAGR (%)", "cagr", True),
        ("Sortino", "sortino", False),
        ("Calmar", "calmar", False),
        ("Avg exposure", "avg_exposure", True),
        ("Cost drag (%)", "cost_drag", True),
        ("n_folds", "n_folds", False),
        ("fold Sharpe mean", "fold_sharpe_mean", False),
        ("fold Sharpe std", "fold_sharpe_std", False),
    ]
    headers = ["Chỉ tiêu"] + [c["label"] for c in columns]
    table: list[list[Any]] = [headers]
    for label, key, as_pct in metric_rows:
        row: list[Any] = [label]
        for col in columns:
            val = col.get(key)
            if val is None:
                row.append("")
            elif key in ("n_trades", "n_folds"):
                row.append(int(val))
            elif as_pct:
                try:
                    row.append(round(float(val) * 100, 4))
                except (TypeError, ValueError):
                    row.append("")
            else:
                try:
                    row.append(round(float(val), 6))
                except (TypeError, ValueError):
                    row.append("")
        table.append(row)

    try:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "report"
        for r_i, row in enumerate(table, start=1):
            for c_i, cell in enumerate(row, start=1):
                ws.cell(r_i, c_i, cell)
        meta_ws = wb.create_sheet("meta")
        meta_ws.append(["key", "value"])
        for k, v in meta.items():
            meta_ws.append(
                [k, json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v]
            )
        folds = meta.get("fold_sharpes") or []
        if folds:
            meta_ws.append([])
            meta_ws.append(["fold_index", "sharpe"])
            for i, s in enumerate(folds, start=1):
                meta_ws.append([i, s])
        xlsx_path = (
            out_path if out_path.suffix.lower() == ".xlsx" else out_path.with_suffix(".xlsx")
        )
        wb.save(xlsx_path)
        return xlsx_path
    except Exception as exc:  # noqa: BLE001
        csv_path = out_path.with_suffix(".csv")
        with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerows(table)
            writer.writerow([])
            writer.writerow(["# export_fallback", str(exc)])
        return csv_path


def main(argv: list[str] | None = None) -> int:
    root = Path(__file__).resolve().parents[1]
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))

    parser = argparse.ArgumentParser(
        description="Backtest presentation report (B0/B1/B2/FW/VNINDEX/VN30)"
    )
    parser.add_argument("--config", default="pipeline/config.yaml")
    parser.add_argument(
        "--universe",
        choices=("smoke", "tickers", "vn100"),
        default="smoke",
        help="smoke=hose_liquid_35; vn100=data/universe/vn100.csv; tickers=--tickers",
    )
    parser.add_argument(
        "--tickers",
        default=None,
        help="Comma-separated (required if --universe tickers; overrides smoke)",
    )
    parser.add_argument(
        "--start-date",
        default=None,
        help="Data/history start (default: oos_start - warmup_years)",
    )
    parser.add_argument("--end-date", default=None)
    parser.add_argument("--oos-start", default=None, help="OOS equity window start")
    parser.add_argument("--oos-end", default=None, help="OOS equity window end")
    parser.add_argument(
        "--warmup-years",
        type=int,
        default=3,
        help="Years of history before OOS for Markov/GARCH/Kalman (default 3)",
    )
    parser.add_argument("--signal-every", type=int, default=1)
    parser.add_argument(
        "--walk-forward",
        dest="walk_forward",
        action="store_true",
        default=False,
        help="Run walk-forward OOS for Framework column",
    )
    parser.add_argument(
        "--no-walk-forward",
        dest="walk_forward",
        action="store_false",
        help="Disable walk-forward (default)",
    )
    parser.add_argument(
        "--with-fundamentals",
        dest="with_fundamentals",
        action="store_true",
        default=True,
        help="Build PIT scoring_schedule (default on for final research)",
    )
    parser.add_argument(
        "--no-fundamentals",
        "--no-with-fundamentals",
        dest="with_fundamentals",
        action="store_false",
        help="Skip fundamental PIT schedule (quant-only / speed)",
    )
    parser.add_argument("--lookback-years", type=int, default=5)
    parser.add_argument("--db-path", default="store/bot.db")
    parser.add_argument("--out-xlsx", default="store/backtest_report.xlsx")
    parser.add_argument("--out-json", default="store/backtest_report.json")
    parser.add_argument(
        "--no-persist-store",
        action="store_true",
        help="Skip writing backtest_results to SQLite",
    )
    parser.add_argument(
        "--refresh-fundamentals",
        action="store_true",
        help="Bỏ qua disk cache scoring_schedule theo năm; fetch lại BCTC",
    )
    parser.add_argument(
        "--refresh-data",
        action="store_true",
        help="Refresh fundamentals cache + bỏ OHLCV disk cache (fetch lại giá)",
    )
    parser.add_argument(
        "--profile",
        action="store_true",
        help="cProfile toàn bộ report; ghi outputs/performance/profile_report.txt",
    )
    args = parser.parse_args(argv)

    if args.profile:
        import cProfile
        import pstats
        from io import StringIO

        profiler = cProfile.Profile()
        profiler.enable()
        try:
            code = _main_impl(args, root)
        finally:
            profiler.disable()
            out = Path("outputs/performance/profile_report.txt")
            out.parent.mkdir(parents=True, exist_ok=True)
            buf = StringIO()
            stats = pstats.Stats(profiler, stream=buf)
            stats.strip_dirs().sort_stats("cumtime").print_stats(30)
            out.write_text(buf.getvalue(), encoding="utf-8")
            _safe_print(f"wrote {out}")
        return code
    return _main_impl(args, root)


def _main_impl(args: argparse.Namespace, root: Path) -> int:
    from copy import deepcopy

    from backtest.ablation import (
        _benchmark_ticker,
        _b1_ta_result,
        _b2_canslim_result,
        _buyhold_result,
        _json_safe,
        _load_config,
        _set_quant_flags,
        _signal_closes,
        build_scoring_schedule,
        load_close_by_ticker,
        persist_ablation_to_store,
        run_ablation,
    )
    from backtest.engine import run_backtest
    from backtest.metrics import compute_metrics
    from backtest.walk_forward import run_walk_forward

    config = _load_config(args.config)
    refresh_fund = bool(args.refresh_fundamentals or args.refresh_data)
    if args.refresh_data:
        # Bỏ OHLCV disk cache cho lần chạy này (provider fetch lại).
        config = deepcopy(config)
        sources = dict(config.get("data_sources") or {})
        price = dict(sources.get("price") or {})
        price["cache_ohlcv"] = False
        sources["price"] = price
        config["data_sources"] = sources
        _safe_print("refresh-data: OHLCV cache disabled for this run")

    oos_end = args.oos_end or args.end_date or date.today().isoformat()
    if args.oos_start:
        oos_start = args.oos_start
    else:
        y, m, d = int(oos_end[:4]), int(oos_end[5:7]), int(oos_end[8:10])
        m2 = m - 6
        y2 = y
        if m2 <= 0:
            m2 += 12
            y2 -= 1
        try:
            oos_start = date(y2, m2, d).isoformat()
        except ValueError:
            oos_start = date(y2, m2, 28).isoformat()

    hist_start = args.start_date or _shift_years(oos_start, -max(int(args.warmup_years), 0))
    hist_end = args.end_date or oos_end
    use_wf = bool(args.walk_forward)

    if args.tickers:
        tickers = [t.strip().upper() for t in str(args.tickers).split(",") if t.strip()]
    elif args.universe == "vn100":
        tickers = _load_vn100_tickers(config, root)
    elif args.universe == "smoke":
        tickers = _load_smoke_tickers(config, root)
    else:
        raise SystemExit("--universe tickers requires --tickers")

    if not tickers:
        raise SystemExit("Empty ticker list")

    _safe_print(
        f"report: n_tickers={len(tickers)} universe={args.universe} "
        f"history={hist_start}..{hist_end} oos={oos_start}..{oos_end} "
        f"signal_every={args.signal_every} walk_forward={use_wf} "
        f"fundamentals={args.with_fundamentals} warmup_years={args.warmup_years}"
    )
    _safe_print(
        "CAVEAT: fixed/current VN100 research universe — subject to survivorship bias; "
        "assumed fundamental publication lag 90d if no filed_at."
    )

    fetch_list = list(dict.fromkeys([*tickers, "VNINDEX", "VN30"]))
    closes = load_close_by_ticker(
        fetch_list,
        hist_start,
        hist_end,
        config,
        include_benchmark=True,
    )
    if not closes:
        raise SystemExit("No OHLCV loaded — check providers / network / cache")

    signal_closes = _signal_closes(
        {t: closes[t] for t in tickers if t in closes},
        config,
    )
    _safe_print(
        f"loaded closes={len(closes)} signal_universe={len(signal_closes)} "
        f"benchmark={_benchmark_ticker(config)} "
        f"has_VNINDEX={'VNINDEX' in closes} has_VN30={'VN30' in closes}"
    )
    data_start = None
    try:
        import pandas as _pd

        first_dates = []
        for _s in closes.values():
            idx = _pd.Index(_pd.to_datetime(_s.index, errors="coerce")).dropna()
            if len(idx):
                first_dates.append(idx.min())
        if first_dates:
            data_start = min(first_dates).date().isoformat()
            if data_start > hist_start:
                _safe_print(
                    f"NOTE: OHLCV earliest={data_start} > requested start={hist_start} "
                    f"(tier limit / cache)."
                )
    except Exception:  # noqa: BLE001
        pass

    scoring_schedule = None
    if args.with_fundamentals:
        _safe_print(
            f"building scoring_schedule (PIT) refresh={refresh_fund}..."
        )
        scoring_schedule = build_scoring_schedule(
            list(signal_closes),
            hist_start,
            hist_end,
            config,
            lookback_years=max(int(args.lookback_years), 1),
            db_path=args.db_path,
            refresh=refresh_fund,
        )
        _safe_print(f"scoring_schedule keys={sorted(scoring_schedule)}")

    _safe_print("baseline B0 buy&hold (OOS)...")
    b0 = _buyhold_result(signal_closes, oos_start, oos_end, config)
    _safe_print("baseline B1 TA (OOS)...")
    b1 = _b1_ta_result(signal_closes, oos_start, oos_end, config)
    _safe_print("baseline B2 CANSLIM (OOS)...")
    b2 = _b2_canslim_result(signal_closes, oos_start, oos_end, config)

    _safe_print("framework stack (regime+alpha+risk) with warmup history...")
    fw_cfg = _set_quant_flags(config, regime=True, alpha=True, risk=True)
    fw_full_raw = run_backtest(
        fw_cfg,
        hist_start,
        hist_end,
        close_by_ticker=closes,
        scoring_schedule=scoring_schedule,
        signal_every_n_days=max(int(args.signal_every), 1),
        signal_tickers=list(signal_closes),
    )
    fw_oos = _slice_equity_to_oos(fw_full_raw, oos_start, oos_end)
    fw_oos["metrics"] = compute_metrics(
        fw_oos["equity_curve"],
        fw_oos["trades"],
        fw_cfg,
        signals=fw_full_raw.get("signals"),
        force_research_metrics=True,
    )
    fw_full = fw_oos

    fold_meta: dict[str, Any] = {}
    fw_col_source = fw_full
    if use_wf:
        _safe_print("walk_forward OOS (framework)...")
        wf = run_walk_forward(
            fw_cfg,
            hist_start,
            hist_end,
            close_by_ticker=closes,
            scoring_schedule=scoring_schedule,
            signal_every_n_days=max(int(args.signal_every), 1),
            signal_tickers=list(signal_closes),
        )
        fold_meta = {
            "n_folds": wf.get("n_folds"),
            "fold_sharpes": wf.get("fold_sharpes") or [],
            "fold_sharpe_mean": wf.get("fold_sharpe_mean"),
            "fold_sharpe_std": wf.get("fold_sharpe_std"),
        }
        fw_col_source = {
            "equity_curve": wf.get("equity_curve") or [],
            "trades": wf.get("trades") or [],
            "metrics": wf.get("metrics") or {},
        }
        _safe_print(
            f"WF folds={fold_meta['n_folds']} "
            f"sharpe_mean={_fmt(fold_meta['fold_sharpe_mean'])} "
            f"sharpe_std={_fmt(fold_meta['fold_sharpe_std'])}"
        )

    _safe_print("ablation layers (OOS window)...")
    ablation = run_ablation(
        config,
        close_by_ticker=closes,
        start_date=oos_start,
        end_date=oos_end,
        scoring_schedule=scoring_schedule,
        signal_every_n_days=max(int(args.signal_every), 1),
    )

    _safe_print("benchmark VNINDEX / VN30 buy&hold (OOS)...")
    vnindex = (
        _buyhold_result({"VNINDEX": closes["VNINDEX"]}, oos_start, oos_end, config)
        if "VNINDEX" in closes
        else None
    )
    vn30 = (
        _buyhold_result({"VN30": closes["VN30"]}, oos_start, oos_end, config)
        if "VN30" in closes
        else None
    )

    columns = [
        _column_from_result("B0", b0),
        _column_from_result("B1", b1),
        _column_from_result("B2", b2),
        _column_from_result(
            "Framework (WF OOS)" if use_wf else "Framework",
            fw_col_source,
            fold_meta=fold_meta or None,
        ),
        _column_from_result("VN-Index", vnindex, is_benchmark=True),
        _column_from_result("VN30", vn30, is_benchmark=True),
    ]

    _safe_print("")
    _safe_print("=== BACKTEST REPORT ===")
    _print_table(columns)
    _safe_print("")
    mfw = dict(fw_col_source.get("metrics") or {})
    _safe_print(
        f"Regime Sharpe: bull={mfw.get('sharpe_bull_display', 'N/A')} "
        f"bear={mfw.get('sharpe_bear_display', 'N/A')} "
        f"neutral={mfw.get('sharpe_neutral_display', 'N/A')}"
    )

    meta = {
        "start_date": hist_start,
        "end_date": hist_end,
        "oos_start": oos_start,
        "oos_end": oos_end,
        "data_start_actual": data_start,
        "n_tickers": len(tickers),
        "tickers": tickers,
        "universe": args.universe,
        "signal_every": int(args.signal_every),
        "walk_forward": use_wf,
        "with_fundamentals": bool(args.with_fundamentals),
        "warmup_years": int(args.warmup_years),
        "execution": "Close T+1",
        "risk_sizing": "inverse_vol_normalize then w_max cap",
        "survivorship": "fixed/current VN100 — subject to survivorship bias",
        "pit_lag_days": 90,
        "fold_sharpes": fold_meta.get("fold_sharpes") or [],
        "fold_sharpe_mean": fold_meta.get("fold_sharpe_mean"),
        "fold_sharpe_std": fold_meta.get("fold_sharpe_std"),
        "n_folds": fold_meta.get("n_folds"),
        "ablation_steps": [
            {
                "layer": s.get("layer"),
                "sharpe": _json_safe(s.get("sharpe")),
                "cagr": _json_safe(s.get("cagr")),
                "max_drawdown": _json_safe(s.get("max_drawdown")),
                "n_trades": s.get("n_trades"),
            }
            for s in (ablation.get("steps") or [])
        ],
        "note": (
            "Khong cook OOS; denser Sharpe -0.84 giu trong DECISIONS. "
            "Execution Close T+1; risk = inverse-vol normalize + w_max."
        ),
    }

    out_xlsx = _export_xlsx_or_csv(columns, Path(args.out_xlsx), meta)
    _safe_print(f"wrote {out_xlsx}")

    payload = {
        **meta,
        "columns": [
            {k: _json_safe(v) if k != "label" else v for k, v in c.items()}
            for c in columns
        ],
        "framework_full_window": {
            "sharpe": _json_safe((fw_full.get("metrics") or {}).get("sharpe")),
            "cagr": _json_safe((fw_full.get("metrics") or {}).get("cagr")),
            "total_return": _json_safe((fw_full.get("metrics") or {}).get("total_return")),
            "n_trades": (fw_full.get("metrics") or {}).get("n_trades"),
            "avg_exposure": _json_safe((fw_full.get("metrics") or {}).get("avg_exposure")),
            "sharpe_bull_display": (fw_full.get("metrics") or {}).get("sharpe_bull_display"),
            "sharpe_bear_display": (fw_full.get("metrics") or {}).get("sharpe_bear_display"),
        },
    }
    out_json = Path(args.out_json)
    out_json.parent.mkdir(parents=True, exist_ok=True)
    out_json.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    _safe_print(f"wrote {out_json}")

    if not args.no_persist_store:
        store_payload = {
            "start_date": oos_start,
            "end_date": oos_end,
            "tickers": tickers,
            "steps": [
                {
                    "layer": "B0_buyhold",
                    "cagr": _json_safe((b0.get("metrics") or {}).get("cagr")),
                    "sharpe": _json_safe((b0.get("metrics") or {}).get("sharpe")),
                    "max_drawdown": _json_safe((b0.get("metrics") or {}).get("max_drawdown")),
                    "win_rate": _json_safe((b0.get("metrics") or {}).get("win_rate")),
                    "n_trades": (b0.get("metrics") or {}).get("n_trades"),
                    "equity_curve": list(b0.get("equity_curve") or []),
                },
                {
                    "layer": "B1_ta",
                    "cagr": _json_safe((b1.get("metrics") or {}).get("cagr")),
                    "sharpe": _json_safe((b1.get("metrics") or {}).get("sharpe")),
                    "max_drawdown": _json_safe((b1.get("metrics") or {}).get("max_drawdown")),
                    "win_rate": _json_safe((b1.get("metrics") or {}).get("win_rate")),
                    "n_trades": (b1.get("metrics") or {}).get("n_trades"),
                    "equity_curve": list(b1.get("equity_curve") or []),
                },
                {
                    "layer": "B2_canslim",
                    "cagr": _json_safe((b2.get("metrics") or {}).get("cagr")),
                    "sharpe": _json_safe((b2.get("metrics") or {}).get("sharpe")),
                    "max_drawdown": _json_safe((b2.get("metrics") or {}).get("max_drawdown")),
                    "win_rate": _json_safe((b2.get("metrics") or {}).get("win_rate")),
                    "n_trades": (b2.get("metrics") or {}).get("n_trades"),
                    "equity_curve": list(b2.get("equity_curve") or []),
                },
                {
                    "layer": "risk",
                    "cagr": _json_safe((fw_full.get("metrics") or {}).get("cagr")),
                    "sharpe": _json_safe((fw_full.get("metrics") or {}).get("sharpe")),
                    "max_drawdown": _json_safe(
                        (fw_full.get("metrics") or {}).get("max_drawdown")
                    ),
                    "win_rate": _json_safe((fw_full.get("metrics") or {}).get("win_rate")),
                    "n_trades": (fw_full.get("metrics") or {}).get("n_trades"),
                    "equity_curve": list(fw_full.get("equity_curve") or []),
                },
            ],
            "walk_forward": {
                "n_folds": fold_meta.get("n_folds"),
                "sharpe": _json_safe((fw_col_source.get("metrics") or {}).get("sharpe")),
                "cagr": _json_safe((fw_col_source.get("metrics") or {}).get("cagr")),
                "n_trades": (fw_col_source.get("metrics") or {}).get("n_trades"),
                "fold_sharpe_mean": _json_safe(fold_meta.get("fold_sharpe_mean")),
                "fold_sharpe_std": _json_safe(fold_meta.get("fold_sharpe_std")),
                "fold_sharpes": [
                    _json_safe(x) for x in (fold_meta.get("fold_sharpes") or [])
                ],
                "equity_curve": list(fw_col_source.get("equity_curve") or []),
            }
            if use_wf
            else None,
        }
        persisted = persist_ablation_to_store(
            store_payload,
            db_path=args.db_path,
            run_id=out_json.stem,
            scope="portfolio",
        )
        _safe_print(f"store: {persisted.get('note')}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
